#!/usr/bin/env python3
"""Excel（パラメータDB）→ data/*.json 変換スクリプト.

使い方:
    python3 tools/export_excel.py                     # リポジトリ直下の最新 .xlsx を自動検出
    python3 tools/export_excel.py path/to/book.xlsx   # ファイルを明示
    python3 tools/export_excel.py --check             # 書き込まず検証のみ

出力:
    data/model.json             回答尺度・7→15翻訳係数・マッチング設定・軸定義
    data/questions.json         30問の質問文と Loading
    data/profiles.json          全プロファイルの15軸・性格座標・結果コメント
    data/roster.json            現行100体のロスター
    data/archetypes.json        アーキタイプ別の15軸リファレンス
    tests/fixtures/excel_baseline.json  Excel Simulator の計算結果（Web実装の突き合わせ用）

このスクリプトは data/comments.json, data/personality-types.json, data/display.json を
上書きしません（Excelに存在しない、手で調整する表現系データのため）。
"""

from __future__ import annotations

import argparse
import datetime as dt
import glob
import json
import math
import re
import os
import sys
import warnings
from collections import Counter, OrderedDict

try:
    import openpyxl
    from openpyxl.utils import column_index_from_string, get_column_letter
except ImportError:  # pragma: no cover - 実行環境向けの案内
    sys.exit("openpyxl が必要です:  pip install openpyxl")

warnings.filterwarnings("ignore", module="openpyxl")

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(REPO_ROOT, "data")
FIXTURE_DIR = os.path.join(REPO_ROOT, "tests", "fixtures")

# ---------------------------------------------------------------------------
# 軸の定義（Excelの列順とコードの軸順を一致させるための唯一の情報源）
# ---------------------------------------------------------------------------

PERSONALITY_AXES = [
    "Initiative",
    "RiskTolerance",
    "Cooperation",
    "SelfReliance",
    "Planning",
    "Adaptability",
    "Mastery",
]

GAME_AXES = [
    "Engage",
    "Dive",
    "Peel",
    "Frontline",
    "Poke",
    "FightTempo",
    "Control",
    "Support",
    "Mobility",
    "SelfSufficiency",
    "FarmScaling",
    "ObjectiveSecure",
    "ScoringSide",
    "Execution",
    "Decision",
]

# Profiles_140 の P_* 列（Excel上は数式）は、各セルの式を読み取って再計算する。
# 例: =ROUND(MAX(0,MIN(100,0.45*K2+0.20*L2+0.15*P2+0.20*Y2)),0)
# 「係数*列」の項を拾い、列見出し（Engage等）へ読み替える。(100-X2) の形は反転項として扱う。
P_FORMULA_TERM = re.compile(r"(-?[0-9.]+)\*(\(100-)?\$?([A-Z]{1,3})\$?\d+")

# Model_Formula の GlobalSetting ブロック。Excelに無い設定は既定値で補い、警告を出す。
GLOBAL_SETTING_DEFAULTS = {
    "PersonalityScale": 50,
    "TranslationScale": 50,
    "PreferenceScale": 50,
    "BaseShapeWeight": 0,
    "SignalReference": 18,
    "SpecificityCorrection": 0,
    # Match_140 の式は Model_Formula!F11 を上限として参照するが、セルが空のことがある。
    # 2026-09-29版のキャッシュ値は ±5 で頭打ちになっているため、既定値を5とする。
    "SpecificityCap": 5,
}

# Model_Formula の DisplaySetting ブロック（Percentile Matching の表示・順位設定）
DISPLAY_SETTING_DEFAULTS = {
    "DisplayBase": 42,
    "DisplayRange": 55,
    "DisplayPower": 3,
    "RankPercentileWeight": 0.8,
    "DisplayMin": 55,
}

# Model_Formula の ResponseStyle ブロック（回答スタイル補正。Simulator の N〜P列で同じ計算をしている）
#   ResponseExtremity = 平均|正規化回答|、補正倍率 = clamp((Target/Extremity)^Strength, MinScale, MaxScale)
#   Extremity が RampExtremity 未満なら強める量を比例して小さくする（全部3は補正なし）
RESPONSE_STYLE_DEFAULTS = {
    "ResponseStyleEnabled": 0,
    "ResponseTarget": 0.425,
    "ResponseStrength": 1,
    "ResponseMinScale": 1,
    "ResponseMaxScale": 1,
    "ResponseRampExtremity": 0,
}

# Interaction_Model の Mode（a=FactorA, b=FactorB の正規化値 -1〜+1）
#   high_high : max(0,a) × max(0,b)   … 両方高いほど効く
#   low_second: max(0,a) × max(0,-b)  … a が高く b が低いほど効く
#   low_first : max(0,-a) × max(0,b)
#   low_low   : max(0,-a) × max(0,-b)
INTERACTION_MODES = {"high_high", "low_second", "low_first", "low_low"}

# 変換中の注意（書き出しは止めない）
NOTICES = []

OFFICIAL_ROLES = ["アタック型", "バランス型", "スピード型", "ディフェンス型", "サポート型"]

# StyleKeywords（Profiles_140）は Axes の「キーワード0側 / キーワード100側」の言葉だけを使い、
# そのプロファイルの軸の値が 50 からこの距離以上、同じ側に寄っていなければならない
KEYWORD_MIN_DISTANCE = 15


# ---------------------------------------------------------------------------
# 汎用ヘルパー
# ---------------------------------------------------------------------------


def excel_round(value: float, digits: int = 0) -> float:
    """ExcelのROUND（四捨五入、0.5は絶対値が大きい方へ）を再現する。"""
    factor = 10 ** digits
    scaled = value * factor
    rounded = math.floor(abs(scaled) + 0.5) * (1 if scaled >= 0 else -1)
    return rounded / factor if digits else int(rounded)


def clean(value):
    """空白のみのセルを None に寄せる。"""
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def as_number(value, default=0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return default


def find_header_row(ws, first_label: str, max_row: int = 30) -> int:
    """A列が first_label の行番号を返す（ヘッダー行の位置ズレに耐えるため）。"""
    for row in range(1, min(ws.max_row, max_row) + 1):
        if clean(ws.cell(row=row, column=1).value) == first_label:
            return row
    raise ValueError(f"'{first_label}' で始まるヘッダー行が {ws.title} に見つかりません")


def header_map(ws, header_row: int) -> "OrderedDict[str, int]":
    """ヘッダー名 → 列番号 の対応を返す。"""
    mapping: "OrderedDict[str, int]" = OrderedDict()
    for col in range(1, ws.max_column + 1):
        name = clean(ws.cell(row=header_row, column=col).value)
        if name and name not in mapping:
            mapping[name] = col
    return mapping


def find_cell(ws, label: str, max_row: int = 60, max_col: int = 30):
    """指定ラベルのセル座標 (row, col) を返す。ブロック位置検出用。"""
    for row in range(1, min(ws.max_row, max_row) + 1):
        for col in range(1, min(ws.max_column, max_col) + 1):
            if clean(ws.cell(row=row, column=col).value) == label:
                return row, col
    raise ValueError(f"'{label}' が {ws.title} に見つかりません")


def strip_zeros(loadings: dict) -> dict:
    """0 の Loading は保存しない（JSONを読みやすく保つ。欠損は0として扱う）。"""
    return {key: value for key, value in loadings.items() if value}


# ---------------------------------------------------------------------------
# 各シートの読み取り
# ---------------------------------------------------------------------------


def read_axes(book) -> tuple:
    """Axes（15プレイ軸）と Personality_Axes（7気質）から、表示名・両端ラベル・定義を読む。

    どちらも1行目が見出し（Code / 表示名 / 区分 / 0側 / 100側 / 定義・意味）。
    定義列に見出しが無い版もあるため、見出しが無ければ「100側」の次の列を定義として扱う。
    Axes には特徴キーワードの語彙（キーワード0側 / キーワード100側）もある。
    """

    def read_sheet(ws, keys):
        cols = header_map(ws, 1)
        definition_col = cols.get("定義") or cols.get("意味") or (cols["100側"] + 1)
        info = {}
        for row in range(2, ws.max_row + 1):
            code = clean(ws.cell(row=row, column=cols["Code"]).value)
            if code not in keys:
                continue
            entry = {
                "key": code,
                "nameJa": clean(ws.cell(row=row, column=cols["表示名"]).value) or code,
                "lowLabel": clean(ws.cell(row=row, column=cols["0側"]).value) or "",
                "highLabel": clean(ws.cell(row=row, column=cols["100側"]).value) or "",
                "description": clean(ws.cell(row=row, column=definition_col).value) or "",
            }
            if "区分" in cols:
                entry["category"] = clean(ws.cell(row=row, column=cols["区分"]).value) or ""
            # 結果カードの特徴キーワードの語彙（15プレイ軸のみ。片側だけの軸もある）
            if "キーワード0側" in cols:
                entry["keywordLow"] = clean(ws.cell(row=row, column=cols["キーワード0側"]).value) or ""
            if "キーワード100側" in cols:
                entry["keywordHigh"] = clean(ws.cell(row=row, column=cols["キーワード100側"]).value) or ""
            info[code] = entry
        missing = [key for key in keys if key not in info]
        if missing:
            raise ValueError(f"{ws.title} シートに定義が無い軸があります: {missing}")
        return info

    return read_sheet(book["Axes"], GAME_AXES), read_sheet(book["Personality_Axes"], PERSONALITY_AXES)


def read_setting_block(ws, label: str, defaults: dict) -> dict:
    """「GlobalSetting | Value」のような縦並びの設定ブロックを読む。無い項目は既定値＋警告。"""
    g_row, g_col = find_cell(ws, label)
    settings = {}
    row = g_row + 1
    while clean(ws.cell(row=row, column=g_col).value):
        settings[clean(ws.cell(row=row, column=g_col).value)] = ws.cell(row=row, column=g_col + 1).value
        row += 1
    for key, default in defaults.items():
        if not isinstance(settings.get(key), (int, float)):
            NOTICES.append(
                f"Model_Formula の {label} に {key} が無いため既定値 {default} を使います"
                f"（{get_column_letter(g_col)}{row}:{get_column_letter(g_col + 1)}{row} に追加すると解消します）"
            )
            settings[key] = default
            row += 1
    return settings


def read_interactions(book) -> list:
    """Interaction_Model シート: 性格2軸の組み合わせで15軸に足すボーナス。"""
    if "Interaction_Model" not in book.sheetnames:
        return []
    ws = book["Interaction_Model"]
    cols = header_map(ws, 1)
    interactions = []
    for row in range(2, ws.max_row + 1):
        name = clean(ws.cell(row=row, column=cols["Interaction"]).value)
        if not name:
            continue
        mode = clean(ws.cell(row=row, column=cols["Mode"]).value)
        if mode not in INTERACTION_MODES:
            raise ValueError(f"Interaction_Model {name}: 未対応の Mode '{mode}'（対応: {sorted(INTERACTION_MODES)}）")
        interactions.append(
            {
                "id": name,
                "factorA": clean(ws.cell(row=row, column=cols["FactorA"]).value),
                "factorB": clean(ws.cell(row=row, column=cols["FactorB"]).value),
                "mode": mode,
                "description": clean(ws.cell(row=row, column=cols["説明"]).value) or "",
                "bonus": strip_zeros(
                    {axis: as_number(ws.cell(row=row, column=cols[axis]).value) for axis in GAME_AXES if axis in cols}
                ),
            }
        )
    return interactions


def read_percentiles(book) -> dict:
    """Percentile_Calibration シート: プロファイルごとの RawMatch 分布（Q00〜Q100 の分位点）。"""
    if "Percentile_Calibration" not in book.sheetnames:
        return {}
    ws = book["__cached__"]["Percentile_Calibration"]
    cols = header_map(ws, 1)
    quantile_cols = [col for name, col in cols.items() if re.fullmatch(r"Q\d{2,3}", name)]
    quantile_cols.sort(key=lambda col: int(clean(ws.cell(row=1, column=col).value)[1:]))
    table = {}
    for row in range(2, ws.max_row + 1):
        profile_id = clean(ws.cell(row=row, column=cols["ProfileID"]).value)
        if not profile_id:
            continue
        table[profile_id] = {
            "rawMean": round(as_number(ws.cell(row=row, column=cols["RawMean"]).value), 4),
            "rawSd": round(as_number(ws.cell(row=row, column=cols["RawSD"]).value), 4),
            "quantiles": [round(as_number(ws.cell(row=row, column=col).value), 4) for col in quantile_cols],
        }
    return table


def read_model(book, game_axis_info, personality_axis_info) -> dict:
    """Model_Formula シートから回答尺度・翻訳係数・マッチング設定を読む。"""
    ws = book["Model_Formula"]

    # 回答尺度（1〜5 → 正規化係数 + 表示ラベル）
    scale_header = find_header_row(ws, "回答値")
    scale = []
    row = scale_header + 1
    while clean(ws.cell(row=row, column=1).value) is not None:
        scale.append(
            {
                "value": int(as_number(ws.cell(row=row, column=1).value)),
                "norm": as_number(ws.cell(row=row, column=2).value),
                "label": clean(ws.cell(row=row, column=3).value) or "",
            }
        )
        row += 1

    # 7性格軸 → 15ゲーム軸 翻訳係数（A列が GameAxis のブロック）
    t_row, t_col = find_cell(ws, "GameAxis")
    t_headers = {
        clean(ws.cell(row=t_row, column=col).value): col
        for col in range(t_col, t_col + len(PERSONALITY_AXES) + 2)
    }
    translation = {}
    row = t_row + 1
    while True:
        axis = clean(ws.cell(row=row, column=t_col).value)
        if not axis:
            break
        translation[axis] = {
            p_axis: as_number(ws.cell(row=row, column=t_headers[p_axis]).value)
            for p_axis in PERSONALITY_AXES
            if p_axis in t_headers
        }
        row += 1

    # 15軸マッチング設定（Weight / OverReqPenalty / PreferenceBlend）
    m_row, m_col = find_cell(ws, "15軸マッチング設定")
    m_header_row = m_row + 1
    m_headers = {
        clean(ws.cell(row=m_header_row, column=col).value): col
        for col in range(m_col, min(ws.max_column, m_col + 6) + 1)
        if clean(ws.cell(row=m_header_row, column=col).value)
    }
    matching = {}
    row = m_header_row + 1
    while True:
        axis = clean(ws.cell(row=row, column=m_headers["GameAxis"]).value)
        if not axis:
            break
        matching[axis] = {
            "weight": as_number(ws.cell(row=row, column=m_headers["Weight"]).value, 1.0),
            "overReqPenalty": as_number(
                ws.cell(row=row, column=m_headers["OverReqPenalty"]).value, 0.0
            ),
            "preferenceBlend": as_number(
                ws.cell(row=row, column=m_headers["PreferenceBlend"]).value, 0.0
            ),
        }
        row += 1

    # 全体設定（各スケール・Shape/Specificity補正）と、表示・順位の設定
    settings = read_setting_block(ws, "GlobalSetting", GLOBAL_SETTING_DEFAULTS)
    display = read_setting_block(ws, "DisplaySetting", DISPLAY_SETTING_DEFAULTS)
    response = read_setting_block(ws, "ResponseStyle", RESPONSE_STYLE_DEFAULTS)

    # 数式の説明（READMEやデバッグ表示で参照する用）
    notes_row = find_header_row(ws, "項目", max_row=60)
    notes = []
    row = notes_row + 1
    while clean(ws.cell(row=row, column=1).value):
        notes.append(
            {
                "item": clean(ws.cell(row=row, column=1).value),
                "formula": clean(ws.cell(row=row, column=2).value) or "",
                "purpose": clean(ws.cell(row=row, column=3).value) or "",
            }
        )
        row += 1

    return {
        "answerScale": scale,
        "neutralAnswer": next((s["value"] for s in scale if s["norm"] == 0), 3),
        "scoreCenter": 50,
        "scoreSpan": 50,
        "scoreMin": 0,
        "scoreMax": 100,
        "scales": {
            "personality": as_number(settings["PersonalityScale"]),
            "translation": as_number(settings["TranslationScale"]),
            "preference": as_number(settings["PreferenceScale"]),
        },
        "matchModel": {
            "baseShapeWeight": as_number(settings["BaseShapeWeight"]),
            "signalReference": as_number(settings["SignalReference"]),
            "specificityCorrection": as_number(settings["SpecificityCorrection"]),
            "specificityCap": as_number(settings["SpecificityCap"]),
            "displayBase": as_number(display["DisplayBase"]),
            "displayRange": as_number(display["DisplayRange"]),
            "displayPower": as_number(display["DisplayPower"]),
            "displayMin": as_number(display["DisplayMin"]),
            "rankPercentileWeight": as_number(display["RankPercentileWeight"]),
        },
        "responseStyle": {
            "enabled": bool(as_number(response["ResponseStyleEnabled"])),
            "target": as_number(response["ResponseTarget"]),
            "strength": as_number(response["ResponseStrength"]),
            "minScale": as_number(response["ResponseMinScale"]),
            "maxScale": as_number(response["ResponseMaxScale"]),
            "rampExtremity": as_number(response["ResponseRampExtremity"]),
        },
        "interactions": read_interactions(book),
        "personalityAxes": [personality_axis_info[k] for k in PERSONALITY_AXES],
        "gameAxes": [game_axis_info[k] for k in GAME_AXES],
        "translation": {axis: translation[axis] for axis in GAME_AXES if axis in translation},
        "matching": {axis: matching[axis] for axis in GAME_AXES if axis in matching},
        "formulaNotes": notes,
    }


def read_questions(book) -> dict:
    """Questions_30 シートから質問文と Loading を読む。"""
    ws = book["Questions_30"]
    header_row = find_header_row(ws, "QID")
    cols = header_map(ws, header_row)
    questions = []
    row = header_row + 1
    while True:
        qid = clean(ws.cell(row=row, column=cols["QID"]).value)
        if not qid:
            break
        personality = {
            axis: as_number(ws.cell(row=row, column=cols[axis]).value)
            for axis in PERSONALITY_AXES
            if axis in cols
        }
        gameplay = {
            axis: as_number(ws.cell(row=row, column=cols[axis]).value)
            for axis in GAME_AXES
            if axis in cols
        }
        questions.append(
            {
                "id": qid,
                "section": clean(ws.cell(row=row, column=cols["区分"]).value) or "",
                "text": clean(ws.cell(row=row, column=cols["質問文"]).value) or "",
                "target": clean(ws.cell(row=row, column=cols["主な測定対象"]).value) or "",
                "note": clean(ws.cell(row=row, column=cols["備考"]).value) or "",
                "intent": clean(ws.cell(row=row, column=cols["設計意図"]).value) or "",
                "personality": strip_zeros(personality),
                "gameplay": strip_zeros(gameplay),
            }
        )
        row += 1
    return {"questions": questions}


def personality_from_formula(formula: str, axes: dict, column_names: dict):
    """P_* 列の式を、そのプロファイルの15軸で計算し直す。解釈できない式なら None。"""
    terms = P_FORMULA_TERM.findall(formula or "")
    if not terms:
        return None
    total = 0.0
    for coeff, inverted, letter in terms:
        axis = column_names.get(column_index_from_string(letter))
        if axis not in axes:
            return None
        value = 100 - axes[axis] if inverted else axes[axis]
        total += float(coeff) * value
    if "MAX(0" in formula.replace(" ", ""):
        total = max(0.0, min(100.0, total))
    return excel_round(total)


def split_keywords(value) -> list:
    text = clean(value)
    if not text:
        return []
    return [word.strip() for word in re.split(r"[・/／,、]", text) if word.strip()]


def read_profiles(book) -> tuple:
    """Profiles_140 シートからプロファイルを読む。P_* 列は数式なので再計算する。
    Percentile_Calibration の分位点もプロファイルごとに持たせる。"""
    ws = book["Profiles_140"]
    percentiles = read_percentiles(book)
    cached = book["__cached__"]["Profiles_140"]
    header_row = find_header_row(ws, "Pokemon")
    cols = header_map(ws, header_row)
    column_names = {col: name for name, col in cols.items()}
    profiles, mismatches = [], []
    row = header_row + 1
    while True:
        pokemon = clean(ws.cell(row=row, column=cols["Pokemon"]).value)
        if not pokemon:
            break
        axes = {
            axis: as_number(ws.cell(row=row, column=cols[axis]).value)
            for axis in GAME_AXES
        }
        personality = {}
        for p_axis in PERSONALITY_AXES:
            column = cols.get(f"P_{p_axis}")
            if column is None:
                continue
            raw = ws.cell(row=row, column=column).value
            cached_value = cached.cell(row=row, column=column).value
            if isinstance(raw, str) and raw.startswith("="):
                value = personality_from_formula(raw, axes, column_names)
                if value is None:
                    value = cached_value
                    mismatches.append(f"{pokemon} P_{p_axis}: 式を解釈できないためキャッシュ値を使用 ({raw})")
                elif isinstance(cached_value, (int, float)) and abs(cached_value - value) > 0.51:
                    # Excelのキャッシュ値と突き合わせて、数式変更に気付けるようにする
                    mismatches.append(f"{pokemon} P_{p_axis}: Excel={cached_value} 再計算={value}")
            else:
                value = raw
            personality[p_axis] = value
        profiles.append(
            {
                "id": clean(ws.cell(row=row, column=cols["ProfileID"]).value),
                "pokemon": pokemon,
                "officialRole": clean(ws.cell(row=row, column=cols["OfficialRole"]).value),
                "profileName": clean(ws.cell(row=row, column=cols["ProfileName"]).value),
                "profileLabel": clean(ws.cell(row=row, column=cols["ProfileLabel"]).value),
                "primaryArchetype": clean(ws.cell(row=row, column=cols["PrimaryArchetype"]).value),
                "secondaryArchetype": clean(
                    ws.cell(row=row, column=cols["SecondaryArchetype"]).value
                ),
                "splitReason": clean(ws.cell(row=row, column=cols["SplitReason"]).value),
                "confidence": clean(ws.cell(row=row, column=cols["Confidence"]).value),
                "resultComment": clean(ws.cell(row=row, column=cols["ResultComment"]).value)
                if "ResultComment" in cols
                else None,
                "styleKeywords": split_keywords(ws.cell(row=row, column=cols["StyleKeywords"]).value)
                if "StyleKeywords" in cols
                else [],
                "rankBias": as_number(ws.cell(row=row, column=cols["RankBias"]).value)
                if "RankBias" in cols
                else 0.0,
                "axes": axes,
                "personality": personality,
            }
        )
        profile_id = profiles[-1]["id"]
        if percentiles:
            if profile_id not in percentiles:
                raise ValueError(f"Percentile_Calibration に {profile_id} がありません")
            profiles[-1]["percentile"] = percentiles[profile_id]
        row += 1
    return {"profiles": profiles}, mismatches


def read_roster(book) -> dict:
    """Roster_100 シートから現行ロスターを読む。"""
    ws = book["Roster_100"]
    header_row = find_header_row(ws, "No")
    cols = header_map(ws, header_row)
    pokemon = []
    row = header_row + 1
    while True:
        name = clean(ws.cell(row=row, column=cols["Pokemon"]).value)
        if not name:
            break
        pokemon.append(
            {
                "no": int(as_number(ws.cell(row=row, column=cols["No"]).value)),
                "name": name,
                "officialRole": clean(ws.cell(row=row, column=cols["OfficialRole"]).value),
                "primaryArchetype": clean(
                    ws.cell(row=row, column=cols["WikiPrimaryArchetype"]).value
                ),
                "secondaryArchetype": clean(
                    ws.cell(row=row, column=cols["WikiSecondaryArchetype"]).value
                ),
                "splitProfile": clean(ws.cell(row=row, column=cols["SplitProfile"]).value) == "Yes",
                "profileCount": int(as_number(ws.cell(row=row, column=cols["ProfileCount"]).value, 1)),
                "confidence": clean(ws.cell(row=row, column=cols["Confidence"]).value),
            }
        )
        row += 1
    counts = Counter(p["officialRole"] for p in pokemon)
    roles = [
        {"key": role, "count": counts.get(role, 0)}
        for role in OFFICIAL_ROLES
        if counts.get(role, 0)
    ]
    # Excelに未知のロールが増えても落とさない
    roles += [
        {"key": role, "count": count}
        for role, count in counts.items()
        if role not in OFFICIAL_ROLES
    ]
    return {"roles": roles, "pokemon": pokemon}


def read_archetypes(book) -> dict:
    """Archetypes シートのアーキタイプ別15軸（参考値）を読む。"""
    ws = book["Archetypes"]
    header_row = find_header_row(ws, "Archetype")
    cols = header_map(ws, header_row)
    archetypes = []
    row = header_row + 1
    while True:
        name = clean(ws.cell(row=row, column=cols["Archetype"]).value)
        if not name:
            break
        archetypes.append(
            {
                "key": name,
                "axes": {
                    axis: as_number(ws.cell(row=row, column=cols[axis]).value)
                    for axis in GAME_AXES
                    if axis in cols
                },
            }
        )
        row += 1
    return {"archetypes": archetypes}


def read_baseline(book) -> dict:
    """Simulator / Match_140 のキャッシュ値を、Web実装の回帰テスト用に取り出す。

    Simulator の D列に入っている回答値と、その結果（7性格軸・15軸・上位10件）を
    そのまま保存する。Excelで回答を変えて再保存すれば、テストの期待値も更新される。
    """
    sim_formula = book["Simulator"]
    sim = book["__cached__"]["Simulator"]

    header_row = find_header_row(sim_formula, "QID")
    answers = []
    row = header_row + 1
    while clean(sim_formula.cell(row=row, column=1).value):
        answers.append(int(as_number(sim_formula.cell(row=row, column=4).value, 3)))
        row += 1

    personality = {
        axis: sim.cell(row=header_row + 1, column=7 + index).value
        for index, axis in enumerate(PERSONALITY_AXES)
    }

    # 「最終15ゲーム軸」ブロック（G列にGameAxisヘッダー）
    block_row, block_col = find_cell(sim_formula, "GameAxis", max_row=40)
    translated, preference, final = {}, {}, {}
    row = block_row + 1
    while True:
        axis = clean(sim_formula.cell(row=row, column=block_col).value)
        if not axis or axis not in GAME_AXES:
            break
        translated[axis] = sim.cell(row=row, column=block_col + 2).value
        preference[axis] = sim.cell(row=row, column=block_col + 3).value
        final[axis] = sim.cell(row=row, column=block_col + 5).value
        row += 1

    match = book["__cached__"]["Match_140"]
    match_header = find_header_row(match, "Pokemon", max_row=20)
    mcols = header_map(match, match_header)
    ranked = []
    row = match_header + 1
    while clean(match.cell(row=row, column=1).value):
        entry = {
            "profileId": clean(match.cell(row=row, column=mcols["ProfileID"]).value),
            "pokemon": clean(match.cell(row=row, column=mcols["Pokemon"]).value),
            "profileName": clean(match.cell(row=row, column=mcols["ProfileName"]).value),
            "rank": match.cell(row=row, column=mcols["Rank"]).value,
        }
        # 各スコアの内訳（Web実装の検算に使う）
        for key, header in (
            ("absoluteScore", "AbsoluteScore"),
            ("shapeScore", "ShapeScore"),
            ("specificityCorrection", "SpecificityCorrection"),
            ("rawMatch", "RawMatch"),
            ("percentile", "PercentileFit"),
            ("displayScore", "DisplayScore"),
            ("rankBias", "RankBias"),
            ("rankScore", "RankScore"),
        ):
            if header in mcols:
                entry[key] = match.cell(row=row, column=mcols[header]).value
        ranked.append(entry)
        row += 1
    ranked.sort(key=lambda item: (item["rank"] or 0))

    return {
        "source": "Excel Simulator / Match_140 のキャッシュ値",
        "answers": answers,
        "personality": personality,
        "translated": translated,
        "preference": preference,
        "final": final,
        "profileCount": len(ranked),
        "top": ranked[:10],
    }


# ---------------------------------------------------------------------------
# 出力
# ---------------------------------------------------------------------------


def write_json(path: str, payload: dict, check_only: bool) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    text = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
    # 数値だけの配列（Percentile の分位点など）は1行にまとめて、ファイルを小さく読みやすくする
    text = re.sub(
        r"\[\s+(-?[0-9.e+-]+(?:,\s+-?[0-9.e+-]+)*)\s+\]",
        lambda match: "[" + ", ".join(part.strip() for part in match.group(1).split(",")) + "]",
        text,
    )
    if check_only:
        print(f"  [check] {os.path.relpath(path, REPO_ROOT)} ({len(text):,} bytes)")
        return
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(text)
    print(f"  書き出し: {os.path.relpath(path, REPO_ROOT)} ({len(text):,} bytes)")


def check_style_keywords(profile_list: list, game_axes: list) -> list:
    """StyleKeywords が語彙表（Axes のキーワード列）にあり、軸の向きと一致しているかを確かめる。"""
    vocabulary = {}
    for axis in game_axes:
        for side in ("Low", "High"):
            word = axis.get(f"keyword{side}")
            if word:
                vocabulary[word] = (axis["key"], side.lower())
    if not vocabulary:
        return ["Axes シートに キーワード0側 / キーワード100側 の列がありません"]
    problems = []
    for profile in profile_list:
        words = profile["styleKeywords"]
        if len(words) != 3 or len(set(words)) != 3:
            problems.append(f"{profile['id']} の StyleKeywords が3つでない・重複している: {words}")
        for word in words:
            if word not in vocabulary:
                problems.append(f"{profile['id']} の StyleKeywords「{word}」が Axes のキーワード列にありません")
                continue
            axis, side = vocabulary[word]
            value = profile["axes"][axis]
            ok = value >= 50 + KEYWORD_MIN_DISTANCE if side == "high" else value <= 50 - KEYWORD_MIN_DISTANCE
            if not ok:
                problems.append(f"{profile['id']} の StyleKeywords「{word}」が {axis}={value:g} と合っていません")
    return problems


def default_workbook() -> str:
    candidates = sorted(
        glob.glob(os.path.join(REPO_ROOT, "*.xlsx")),
        key=os.path.getmtime,
        reverse=True,
    )
    candidates = [path for path in candidates if not os.path.basename(path).startswith("~$")]
    if not candidates:
        sys.exit("リポジトリ直下に .xlsx が見つかりません。パスを引数で指定してください。")
    return candidates[0]


class WorkbookPair(dict):
    """数式シートとキャッシュ値シートの両方へ同じ書き方でアクセスするためのラッパー。

    book["Profiles_140"]              → 数式を保持したシート
    book["__cached__"]["Profiles_140"] → 計算結果（キャッシュ値）のシート
    """

    def __init__(self, formulas, cached):
        super().__init__()
        self._formulas = formulas
        self._cached = cached

    def __getitem__(self, key):
        if key == "__cached__":
            return self._cached
        return self._formulas[key]

    @property
    def sheetnames(self):
        return self._formulas.sheetnames


def main() -> int:
    parser = argparse.ArgumentParser(description="Excel パラメータDB → data/*.json 変換")
    parser.add_argument("workbook", nargs="?", help="Excelファイルのパス")
    parser.add_argument("--check", action="store_true", help="書き込まずに検証のみ行う")
    args = parser.parse_args()

    workbook_path = args.workbook or default_workbook()
    print(f"読み込み: {os.path.relpath(workbook_path, REPO_ROOT)}")

    book = WorkbookPair(
        openpyxl.load_workbook(workbook_path, data_only=False),
        openpyxl.load_workbook(workbook_path, data_only=True),
    )
    return export(book, workbook_path, args.check)


def export(book, workbook_path: str, check_only: bool) -> int:
    meta = {
        "generatedAt": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        "sourceWorkbook": os.path.basename(workbook_path),
        "note": "このファイルは tools/export_excel.py が生成します。直接編集せずExcel側を修正してください。",
    }

    game_axis_info, personality_axis_info = read_axes(book)
    model = read_model(book, game_axis_info, personality_axis_info)
    questions = read_questions(book)
    profiles, mismatches = read_profiles(book)
    roster = read_roster(book)
    archetypes = read_archetypes(book)
    baseline = read_baseline(book)

    # --- 検証 ---
    problems = []
    profile_list = profiles["profiles"]
    roster_list = roster["pokemon"]

    if len(set(p["id"] for p in profile_list)) != len(profile_list):
        problems.append("ProfileID が重複しています")
    roster_names = {p["name"] for p in roster_list}
    profile_names = {p["pokemon"] for p in profile_list}
    if profile_names - roster_names:
        problems.append(f"Roster に無いポケモン: {sorted(profile_names - roster_names)}")
    if roster_names - profile_names:
        problems.append(f"Profile が無いポケモン: {sorted(roster_names - profile_names)}")
    for axis in GAME_AXES:
        if axis not in model["translation"]:
            problems.append(f"翻訳係数に {axis} がありません")
        if axis not in model["matching"]:
            problems.append(f"マッチング設定に {axis} がありません")
    for profile in profile_list:
        if not profile.get("resultComment"):
            NOTICES.append(f"{profile['id']} に ResultComment がありません（汎用コメントで代用します）")
        for axis, value in profile["axes"].items():
            if not 0 <= value <= 100:
                problems.append(f"{profile['id']} の {axis} が0〜100の外: {value}")
    problems.extend(check_style_keywords(profile_list, model["gameAxes"]))
    counted = Counter(p["pokemon"] for p in profile_list)
    for entry in roster_list:
        if counted[entry["name"]] != entry["profileCount"]:
            problems.append(
                f"{entry['name']}: ProfileCount={entry['profileCount']} だが"
                f" Profiles_140 上は {counted[entry['name']]} 件"
            )

    print(
        f"件数: ロスター {len(roster_list)}体 / プロファイル {len(profile_list)}件 /"
        f" 質問 {len(questions['questions'])}問 /"
        f" 軸 {len(model['personalityAxes'])}+{len(model['gameAxes'])}"
    )
    for role in roster["roles"]:
        print(f"  {role['key']}: {role['count']}体")

    for message in mismatches[:10]:
        print(f"  [注意] 性格座標がExcelのキャッシュ値と一致しません: {message}")
    if len(mismatches) > 10:
        print(f"  [注意] ほか {len(mismatches) - 10} 件")

    for message in NOTICES:
        print(f"  [注意] {message}")

    if problems:
        print("\n検証エラー:")
        for problem in problems:
            print(f"  - {problem}")
        return 1

    counts = {
        "roster": len(roster_list),
        "profiles": len(profile_list),
        "questions": len(questions["questions"]),
        "personalityAxes": len(model["personalityAxes"]),
        "gameAxes": len(model["gameAxes"]),
    }

    write_json(os.path.join(DATA_DIR, "model.json"), {"meta": meta, **model}, check_only)
    write_json(
        os.path.join(DATA_DIR, "questions.json"),
        {"meta": {**meta, "count": counts["questions"]}, **questions},
        check_only,
    )
    write_json(
        os.path.join(DATA_DIR, "profiles.json"),
        {"meta": {**meta, "count": counts["profiles"]}, **profiles},
        check_only,
    )
    write_json(
        os.path.join(DATA_DIR, "roster.json"),
        {"meta": {**meta, "count": counts["roster"]}, **roster},
        check_only,
    )
    write_json(os.path.join(DATA_DIR, "archetypes.json"), {"meta": meta, **archetypes}, check_only)
    write_json(
        os.path.join(FIXTURE_DIR, "excel_baseline.json"),
        {"meta": meta, "counts": counts, **baseline},
        check_only,
    )
    print("完了" if not check_only else "検証のみ完了（ファイルは書き換えていません）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
