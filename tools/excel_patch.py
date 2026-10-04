#!/usr/bin/env python3
"""パラメータDB（xlsx）を、書式・条件付き書式・数式を壊さずに書き換え、数式のキャッシュ値も再計算して更新する。

    python3 tools/excel_patch.py --spec changes.json            # 値・数式の変更
    python3 tools/excel_patch.py --calibration excel-handoff/calibration.json
                                                                 # Percentile 分位点と RankBias を書き戻す
    python3 tools/excel_patch.py --verify                        # 書き換えずに、キャッシュ値が再計算と一致するかだけ確認

openpyxl で保存すると条件付き書式の拡張などが消えるため、xlsx の XML を直接書き換える。
数式のキャッシュ値は formulas（pip install formulas）で再計算した値に置き換え、最後にもう一度再計算して
すべての数式セルが一致することを確かめる（一致しなければ書き込まずに終了）。

changes.json:
{
  "merges":   [["Model_Formula", "A11:J11", "A11:D11"]],                 # 結合範囲の変更
  "removeMerges": [["Match_140", "A3:X3"]],                              # 結合の解除
  "profiles": [{"id": "ストリンダー-1", "col": "FarmScaling", "value": 35}],  # Profiles_140（ProfileID で行を探す）
  "cells":    [{"sheet": "Model_Formula", "ref": "B14", "value": 0.3}],   # 値（無いセルは作る）
  "formulas": [{"sheet": "Simulator", "ref": "N5", "formula": "MAX(-1,MIN(1,E5*$P$7))"}],
  "replaceInFormulas": [{"sheet": "Simulator", "find": "$E$5:$E$34", "replace": "$N$5:$N$34"}]
}
"""

from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys
import tempfile
import warnings
import zipfile
from xml.sax.saxutils import escape

warnings.simplefilter("ignore")
REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CELL = r'<x:c r="%s"(?P<attrs>(?: [^>]*?)?)(?:/>|>(?P<body>.*?)</x:c>)'


def default_workbook() -> str:
    paths = [p for p in glob.glob(os.path.join(REPO_ROOT, "*.xlsx")) if not os.path.basename(p).startswith("~$")]
    if not paths:
        sys.exit("リポジトリ直下に .xlsx が見つかりません。")
    return max(paths, key=os.path.getmtime)


def column_number(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def split_ref(ref: str) -> tuple[str, int]:
    m = re.match(r"([A-Z]+)(\d+)$", ref)
    return m.group(1), int(m.group(2))


def format_value(value):
    if isinstance(value, bool):
        return "b", "1" if value else "0"
    if isinstance(value, (int, float)):
        return "n", str(int(value)) if float(value).is_integer() else repr(float(value))
    if isinstance(value, str) and value.startswith("#"):
        return "e", escape(value)
    return "str", escape(str(value))


class Workbook:
    def __init__(self, path: str):
        with zipfile.ZipFile(path) as z:
            self.items = [(info, z.read(info.filename)) for info in z.infolist()]
        self.content = {info.filename: data for info, data in self.items}
        book = self.content["xl/workbook.xml"].decode()
        rels = self.content["xl/_rels/workbook.xml.rels"].decode()
        targets = dict((m.group(2), m.group(1)) for m in re.finditer(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels))
        self.paths = {
            m.group(1): targets[m.group(2)].lstrip("/")
            for m in re.finditer(r'<x:sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', book)
        }

    def sheet(self, name: str) -> str:
        return self.content[self.paths[name]].decode()

    def put(self, name: str, text: str) -> None:
        self.content[self.paths[name]] = text.encode()

    def save(self, path: str) -> None:
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
            for info, _ in self.items:
                out.writestr(info, self.content[info.filename])


def find_cell(text: str, ref: str):
    return re.compile(CELL % ref, re.S).search(text)


def upsert_cell(text: str, ref: str, cell_xml: str) -> str:
    """セルがあれば置き換え、無ければ行の中の正しい列位置に挿入する（行が無ければ作る）。"""
    m = find_cell(text, ref)
    if m:
        return text[: m.start()] + cell_xml + text[m.end() :]
    col, row = split_ref(ref)
    row_match = re.search(r'(<x:row r="%d"[^>]*?)(/>|>(.*?)</x:row>)' % row, text, re.S)
    if not row_match:
        # 行を番号順の位置に作る
        rows = [(int(r.group(1)), r.start()) for r in re.finditer(r'<x:row r="(\d+)"', text)]
        after = [start for number, start in rows if number > row]
        position = after[0] if after else text.index("</x:sheetData>")
        return text[:position] + f'<x:row r="{row}">{cell_xml}</x:row>' + text[position:]
    if row_match.group(2) == "/>":
        return text[: row_match.start()] + f"{row_match.group(1)}>{cell_xml}</x:row>" + text[row_match.end() :]
    body_start = row_match.start(3)
    body = row_match.group(3)
    for cell in re.finditer(r'<x:c r="([A-Z]+)\d+"', body):
        if column_number(cell.group(1)) > column_number(col):
            insert_at = body_start + cell.start()
            return text[:insert_at] + cell_xml + text[insert_at:]
    insert_at = body_start + len(body)
    return text[:insert_at] + cell_xml + text[insert_at:]


def keep_style(match) -> str:
    if not match:
        return ""
    style = re.search(r'\ss="\d+"', match.group("attrs") or "")
    return style.group(0) if style else ""


def set_value(text: str, ref: str, value) -> str:
    m = find_cell(text, ref)
    if m and "<x:f>" in (m.group("body") or ""):
        raise SystemExit(f"{ref} は数式セルです（formulas で指定してください）")
    kind, encoded = format_value(value)
    return upsert_cell(text, ref, f'<x:c r="{ref}"{keep_style(m)} t="{kind}"><x:v>{encoded}</x:v></x:c>')


def set_formula(text: str, ref: str, formula: str) -> str:
    m = find_cell(text, ref)
    formula = formula[1:] if formula.startswith("=") else formula
    return upsert_cell(text, ref, f'<x:c r="{ref}"{keep_style(m)} t="n"><x:f>{escape(formula)}</x:f><x:v>0</x:v></x:c>')


def set_cache(text: str, ref: str, value) -> tuple[str, bool]:
    m = find_cell(text, ref)
    if not m or "<x:f>" not in (m.group("body") or ""):
        return text, False
    formula = re.search(r"<x:f>.*?</x:f>", m.group("body"), re.S).group(0)
    kind, encoded = format_value(value)
    new = f'<x:c r="{ref}"{keep_style(m)} t="{kind}">{formula}<x:v>{encoded}</x:v></x:c>'
    return text[: m.start()] + new + text[m.end() :], True


def profile_lookup(text: str):
    header = {m.group(2): m.group(1) for m in re.finditer(r'<x:c r="([A-Z]+)1"[^>]*t="str"><x:v>([^<]+)</x:v>', text)}
    rows = {m.group(2): m.group(1) for m in re.finditer(r'<x:c r="C(\d+)"[^>]*t="str"><x:v>([^<]+)</x:v>', text)}
    return header, rows


def apply_spec(book: Workbook, spec: dict) -> int:
    changed = 0
    for sheet_name, old, new in spec.get("merges", []):
        text = book.sheet(sheet_name)
        tag = f'<x:mergeCell ref="{old}" />'
        if tag not in text:
            raise SystemExit(f"{sheet_name} に {old} の結合がありません")
        book.put(sheet_name, text.replace(tag, f'<x:mergeCell ref="{new}" />'))
        changed += 1
    for sheet_name, ref in spec.get("removeMerges", []):
        text = book.sheet(sheet_name)
        tag = f'<x:mergeCell ref="{ref}" />'
        if tag not in text:
            raise SystemExit(f"{sheet_name} に {ref} の結合がありません")
        text = text.replace(tag, "")
        text = re.sub(r"<x:mergeCells>\s*</x:mergeCells>", "", text)
        book.put(sheet_name, text)
        changed += 1
    if spec.get("profiles"):
        text = book.sheet("Profiles_140")
        header, rows = profile_lookup(text)
        for edit in spec["profiles"]:
            text = set_value(text, f"{header[edit['col']]}{rows[edit['id']]}", edit["value"])
            changed += 1
        book.put("Profiles_140", text)
    for edit in spec.get("cells", []):
        book.put(edit["sheet"], set_value(book.sheet(edit["sheet"]), edit["ref"], edit["value"]))
        changed += 1
    for edit in spec.get("formulas", []):
        book.put(edit["sheet"], set_formula(book.sheet(edit["sheet"]), edit["ref"], edit["formula"]))
        changed += 1
    for edit in spec.get("replaceInFormulas", []):
        text = book.sheet(edit["sheet"])
        find, replace = escape(edit["find"]), escape(edit["replace"])
        count = 0

        def swap(match):
            nonlocal count
            count += match.group(0).count(find)
            return match.group(0).replace(find, replace)

        text = re.sub(r"<x:f>.*?</x:f>", swap, text, flags=re.S)
        if not count:
            raise SystemExit(f"{edit['sheet']} の数式に {edit['find']} がありません")
        book.put(edit["sheet"], text)
        changed += count
    return changed


def calibration_spec(book: Workbook, calibration: dict) -> dict:
    """calibration.json（tools/calibrate-model.mjs の出力）を、Percentile_Calibration と Profiles_140 RankBias の値にする。"""
    text = book.sheet("Percentile_Calibration")
    header = {m.group(2): m.group(1) for m in re.finditer(r'<x:c r="([A-Z]+)1"[^>]*t="str"><x:v>([^<]+)</x:v>', text)}
    rows = {m.group(2): m.group(1) for m in re.finditer(r'<x:c r="A(\d+)"[^>]*t="str"><x:v>([^<]+)</x:v>', text)}
    cells = []
    for profile_id, entry in calibration["percentiles"].items():
        row = rows[profile_id]
        cells.append({"sheet": "Percentile_Calibration", "ref": f"{header['RawMean']}{row}", "value": entry["rawMean"]})
        cells.append({"sheet": "Percentile_Calibration", "ref": f"{header['RawSD']}{row}", "value": entry["rawSd"]})
        for k, value in enumerate(entry["quantiles"]):
            cells.append({"sheet": "Percentile_Calibration", "ref": f"{header[f'Q{k:02d}']}{row}", "value": value})
    profiles = [{"id": pid, "col": "RankBias", "value": value} for pid, value in calibration["rankBias"].items()]
    return {"cells": cells, "profiles": profiles}


def recalculate(path: str) -> dict:
    import formulas

    solution = formulas.ExcelModel().loads(path).finish().calculate()
    values = {}
    for key, cell in solution.items():
        m = re.match(r"'\[(.+)\](.+)'!([A-Z]+\d+)$", key)
        if not m:
            continue
        value = cell.value[0][0] if hasattr(cell, "value") else cell
        if hasattr(value, "item"):
            value = value.item()
        if value is None or str(value) == "empty":
            continue
        values.setdefault(m.group(2), {})[m.group(3)] = value
    return values


def write_caches(book: Workbook, values: dict) -> int:
    upper = {name.upper(): name for name in book.paths}
    updated = 0
    for sheet_upper, cells in values.items():
        name = upper.get(sheet_upper)
        if not name:
            continue
        text = book.sheet(name)
        for ref, value in cells.items():
            text, changed = set_cache(text, ref, value)
            updated += changed
        book.put(name, text)
    return updated


def verify(path: str) -> list:
    """キャッシュ値と再計算値が食い違う数式セルの一覧。"""
    import openpyxl

    values = recalculate(path)
    formulas_book = openpyxl.load_workbook(path)
    cached_book = openpyxl.load_workbook(path, data_only=True)
    upper = {name.upper(): name for name in formulas_book.sheetnames}
    mismatches = []
    for sheet_upper, cells in values.items():
        name = upper.get(sheet_upper)
        if not name:
            continue
        for ref, value in cells.items():
            formula = formulas_book[name][ref].value
            if not (isinstance(formula, str) and formula.startswith("=")):
                continue
            cached = cached_book[name][ref].value
            if isinstance(cached, (int, float)) and not isinstance(cached, bool):
                same = isinstance(value, (int, float)) and abs(float(value) - float(cached)) < 1e-6
            else:
                same = str(value) == str(cached)
            if not same:
                mismatches.append((name, ref, cached, value))
    return mismatches


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", help="xlsx のパス（省略時はリポジトリ直下の最新）")
    parser.add_argument("--spec", help="値・数式の変更（JSON）")
    parser.add_argument("--calibration", help="tools/calibrate-model.mjs が出力した calibration.json")
    parser.add_argument("--verify", action="store_true", help="キャッシュ値が再計算と一致するかだけ確認する")
    args = parser.parse_args()
    path = args.workbook or default_workbook()

    if args.verify:
        mismatches = verify(path)
        print(f"数式セルのキャッシュ値と再計算値の食い違い: {len(mismatches)} 件")
        for item in mismatches[:20]:
            print("  ", item)
        return 1 if mismatches else 0

    book = Workbook(path)
    spec = {}
    if args.spec:
        with open(args.spec, encoding="utf-8") as handle:
            spec = json.load(handle)
    if args.calibration:
        with open(args.calibration, encoding="utf-8") as handle:
            extra = calibration_spec(book, json.load(handle))
        spec = {**spec, "cells": spec.get("cells", []) + extra["cells"], "profiles": spec.get("profiles", []) + extra["profiles"]}
    if not spec:
        parser.error("--spec か --calibration を指定してください")

    changed = apply_spec(book, spec)
    with tempfile.TemporaryDirectory() as tmp:
        work = os.path.join(tmp, os.path.basename(path))
        book.save(work)
        updated = write_caches(book, recalculate(work))
        book.save(work)
        mismatches = verify(work)
        if mismatches:
            print(f"再計算後もキャッシュ値が一致しない数式セルが {len(mismatches)} 件あるため、書き込みを中止しました。")
            for item in mismatches[:20]:
                print("  ", item)
            return 1
        book.save(path)
    print(f"{os.path.relpath(path, REPO_ROOT)}: {changed} 件を変更、数式キャッシュ {updated} 件を再計算値で更新（全数式セル一致を確認済み）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
